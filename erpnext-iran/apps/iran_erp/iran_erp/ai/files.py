"""Files the user attaches to an assistant question: read as tables, or passed to the model."""

import base64
import csv
import datetime
import io
import mimetypes
import re

import frappe
from frappe import _

from iran_erp.utils.jalali import to_gregorian
from iran_erp.utils.national_id import _NON_LATIN_DIGITS

# Images and PDFs larger than this are not sent to the model
MAX_MODEL_FILE_BYTES = 8 * 1024 * 1024
MODEL_FILE_TYPES = {"image/png", "image/jpeg", "image/webp", "image/gif", "application/pdf"}
TABLE_EXTENSIONS = {"csv", "xlsx"}


def get_file(file_url: str):
	"""The File record for a URL, if the current user may read it."""
	name = frappe.db.get_value("File", {"file_url": file_url}, "name")
	if not name:
		frappe.throw(_("File {0} was not found").format(file_url))
	file = frappe.get_doc("File", name)
	file.check_permission("read")
	return file


def content_type(file) -> str:
	return mimetypes.guess_type(file.file_name or file.file_url)[0] or ""


def extension(file) -> str:
	return (file.file_name or file.file_url).rsplit(".", 1)[-1].lower()


def model_part(file) -> dict | None:
	"""A chat message part that lets the model see an image or PDF, or None for other files."""
	mime = content_type(file)
	if mime not in MODEL_FILE_TYPES:
		return None
	content = file.get_content()
	if isinstance(content, str):
		content = content.encode()
	if len(content) > MAX_MODEL_FILE_BYTES:
		return None
	data_url = f"data:{mime};base64,{base64.b64encode(content).decode()}"
	if mime == "application/pdf":
		return {"type": "file", "file": {"filename": file.file_name, "file_data": data_url}}
	return {"type": "image_url", "image_url": {"url": data_url}}


def read_table(file) -> list[list]:
	"""Rows of a CSV or Excel file, as lists of cell values."""
	ext = extension(file)
	if ext not in TABLE_EXTENSIONS:
		frappe.throw(
			_("Only CSV and XLSX files can be read as tables. Save the statement as XLSX and try again.")
		)
	content = file.get_content()
	if ext == "xlsx":
		from openpyxl import load_workbook

		if isinstance(content, str):
			content = content.encode()
		sheet = load_workbook(io.BytesIO(content), read_only=True, data_only=True).worksheets[0]
		rows = [list(row) for row in sheet.iter_rows(values_only=True)]
	else:
		if isinstance(content, bytes):
			# Iranian bank exports are UTF-8 or Windows-1256
			for encoding in ("utf-8-sig", "cp1256"):
				try:
					content = content.decode(encoding)
					break
				except UnicodeDecodeError:
					continue
		rows = list(csv.reader(io.StringIO(content)))
	# Drop rows that are entirely empty
	return [row for row in rows if any(cell not in (None, "") for cell in row)]


def latin_digits(value) -> str:
	return str(value).translate(_NON_LATIN_DIGITS)


def parse_amount(value) -> float | None:
	"""Amounts as banks export them: 1,250,000 / the same in Persian digits / (500) / 500- / 1250000.00"""
	if value in (None, ""):
		return None
	if isinstance(value, int | float):
		return float(value)
	text = latin_digits(value).strip()
	negative = (text.startswith("(") and text.endswith(")")) or text.endswith("-") or text.startswith("-")
	text = re.sub(r"[^\d.]", "", text.replace("٫", "."))
	if not text or text == ".":
		return None
	amount = float(text)
	return -amount if negative else amount


def parse_date(value) -> datetime.date | None:
	"""Dates in Jalali (1405/07/20, 14050720) or Gregorian (2026-10-12) form, or Excel dates."""
	if value in (None, ""):
		return None
	if isinstance(value, datetime.datetime):
		return value.date()
	if isinstance(value, datetime.date):
		return value
	text = latin_digits(value).strip()
	match = re.search(r"(\d{4})[/\-.](\d{1,2})[/\-.](\d{1,2})", text) or re.fullmatch(
		r"(\d{4})(\d{2})(\d{2})", text.split()[0] if text else ""
	)
	if not match:
		return None
	year, month, day = (int(part) for part in match.groups())
	try:
		if year < 1700:
			return to_gregorian(year, month, day)
		return datetime.date(year, month, day)
	except ValueError:
		return None
