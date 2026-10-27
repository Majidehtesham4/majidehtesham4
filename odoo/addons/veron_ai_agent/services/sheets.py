"""Reading a spreadsheet given to an agent: the first rows of each sheet of an xlsx file (values, not
formulas) or of a CSV file; a header row names the columns."""

import csv
import io

MAX_ROWS = 2000
MAX_COLUMNS = 40


def read(data, mimetype, max_rows=MAX_ROWS):
	"""{sheet name: [row, ...]} with each row a list of plain values"""
	if mimetype == "text/csv":
		return {"CSV": _csv(data, max_rows)}
	from openpyxl import load_workbook

	book = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
	sheets = {}
	try:
		for sheet in book.worksheets[:10]:
			rows = []
			for row in sheet.iter_rows(max_row=max_rows, max_col=MAX_COLUMNS, values_only=True):
				if any(cell not in (None, "") for cell in row):
					rows.append([_plain(cell) for cell in row])
			sheets[sheet.title] = _trim(rows)
	finally:
		book.close()
	return sheets


def _csv(data, max_rows):
	for encoding in ("utf-8-sig", "cp1256"):
		try:
			text = data.decode(encoding)
			break
		except UnicodeDecodeError:
			continue
	try:
		dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
	except csv.Error:
		dialect = csv.excel
	rows = []
	for row in csv.reader(io.StringIO(text), dialect):
		if any(cell.strip() for cell in row):
			rows.append([_plain(cell) for cell in row[:MAX_COLUMNS]])
		if len(rows) >= max_rows:
			break
	return _trim(rows)


def _trim(rows):
	"""Without the empty columns at the end, which a sheet often reports"""
	width = max((max((i + 1 for i, cell in enumerate(row) if cell not in (None, "")), default=0) for row in rows), default=0)
	return [list(row[:width]) + [""] * (width - len(row)) for row in rows]


def _plain(value):
	if value is None:
		return ""
	if isinstance(value, float) and value.is_integer():
		return int(value)
	if isinstance(value, (int, float)):
		return value
	if hasattr(value, "isoformat"):
		return value.isoformat()[:10]
	text = str(value).strip()
	if len(text) > 1 and text[0] in "0+۰٠" or len(text) > 15:
		# A phone, a code or an account number written as text stays text
		return text
	digits = text.translate(str.maketrans("۰۱۲۳۴۵۶۷۸۹٠١٢٣٤٥٦٧٨٩", "01234567890123456789")).replace(",", "").replace("٬", "")
	try:
		number = float(digits)
		return int(number) if number.is_integer() else number
	except ValueError:
		return text


def table(rows):
	"""(header, records as dicts) of a sheet whose first row names the columns"""
	if not rows:
		return [], []
	header = [str(cell or f"col{index + 1}").strip() for index, cell in enumerate(rows[0])]
	records = [dict(zip(header, row)) for row in rows[1:]]
	return header, records
